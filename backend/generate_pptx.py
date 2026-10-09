import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # Blank layout

    # Colors
    bg_color = RGBColor(2, 6, 23)        # Slate-950 #020617
    panel_color = RGBColor(15, 23, 42)    # Slate-900 #0F172A
    border_color = RGBColor(30, 41, 59)   # Slate-800 #1E293B
    cyan_text = RGBColor(34, 211, 238)    # Cyan-400 #22D3EE
    purple_text = RGBColor(192, 132, 252) # Purple-400 #C084FC
    white_text = RGBColor(248, 250, 252)  # Slate-50 #F8FAFC
    gray_text = RGBColor(148, 163, 184)   # Slate-400 #94A3B8
    emerald_text = RGBColor(52, 211, 153) # Emerald-400 #34D399
    red_text = RGBColor(248, 113, 113)    # Red-400 #F87171

    # Helper: Set background color of slide
    def set_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color

    # Helper: Add header and footer to slide
    def add_header_footer(slide, title_text, slide_num):
        # Header line
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.8), Inches(12.333), Inches(0.02))
        shape.fill.solid()
        shape.fill.fore_color.rgb = border_color
        shape.line.fill.background()

        # Header tag
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(6), Inches(0.4))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "FEDSHIELD-ID  //  HACKATHON PRODUCT PITCH"
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = gray_text

        # Title
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.95), Inches(12.333), Inches(0.6))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = cyan_text

        # Footer line
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(6.8), Inches(12.333), Inches(0.02))
        shape.fill.solid()
        shape.fill.fore_color.rgb = border_color
        shape.line.fill.background()

        # Footer text
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(8), Inches(0.4))
        tf = tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = "CONFIDENTIAL  |  PRIVACY-FIRST BANKING IDENTITY TRUST"
        p.font.name = "Arial"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = gray_text

        # Slide Number
        txBox = slide.shapes.add_textbox(Inches(10.533), Inches(6.9), Inches(2.3), Inches(0.4))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        p.text = f"Slide {slide_num} of 4"
        p.font.name = "Arial"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = gray_text

    # Helper: Create stylized panel box
    def create_panel(slide, left, top, width, height, title, items, title_color=cyan_text):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = panel_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)

        txBox = slide.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.2), Inches(width - 0.4), Inches(height - 0.4))
        tf = txBox.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = title_color
        p.space_after = Pt(14)

        for heading, desc in items:
            p = tf.add_paragraph()
            p.space_after = Pt(10)
            
            run = p.add_run()
            run.text = f"• {heading}: "
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = white_text

            run2 = p.add_run()
            run2.text = desc
            run2.font.name = "Arial"
            run2.font.size = Pt(11)
            run2.font.color.rgb = gray_text

    # ==================== SLIDE 0: TITLE SLIDE ====================
    slide0 = prs.slides.add_slide(blank_layout)
    set_background(slide0)

    # Title box
    txBox = slide0.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.333), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "FedShield-ID"
    p.font.name = "Arial"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = cyan_text

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Privacy-First Identity Trust Platform for Banking Networks"
    p2.font.name = "Arial"
    p2.font.size = Pt(22)
    p2.font.color.rgb = white_text
    p2.space_before = Pt(10)

    # Metadata box
    txBox = slide0.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11.333), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Hackathon Pitch Deck  |  20–22 July 2026"
    p.font.name = "Arial"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = purple_text

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Topic: Cybersecurity & Privacy-Preserving AI in Financial Tech\nPresenter: Mohd. Amaan Hamid (AI/ML Architect & Team Lead)"
    p2.font.name = "Arial"
    p2.font.size = Pt(12)
    p2.font.color.rgb = gray_text
    p2.space_before = Pt(8)

    # ==================== SLIDE 1: PROBLEM & SOLUTION ====================
    slide1 = prs.slides.add_slide(blank_layout)
    set_background(slide1)
    add_header_footer(slide1, "Slide 1: Problem & Solution", 1)

    prob_items = [
        ("Identity Silos & Centralization", "Legacy architectures pool raw customer PII. This violates banking privacy regulations (GDPR, RBI) and creates data breach liabilities."),
        ("Synthetic ID & Account Takeovers", "Fraudsters bypass static onboarding checks and hijack account recovery channels (e.g. SIM swaps) before transactions occur."),
        ("Robotic Credential Sweeps", "Credential stuffing attacks utilize automated bot scripts, mimicking standard logins but clicking in sub-millisecond rates.")
    ]
    create_panel(slide1, 0.5, 1.7, 5.9, 4.8, "The Challenge in Banking Systems", prob_items, title_color=red_text)

    sol_items = [
        ("Decentralized Machine Learning", "Models are trained locally at individual banking nodes without data pooling. Global models are merged via secure FedAvg federated rounds."),
        ("Continuous Multi-Dimension Scoring", "Aggregates 10 scoring parameters (reputation, location, biometrics, insider activity) into an active trust score (0–100)."),
        ("Risk-Based Step-Up Challenges", "Triggers dynamic security checks (like Biometric Face ID) only when scoring parameters indicate risk anomalies.")
    ]
    create_panel(slide1, 6.9, 1.7, 5.9, 4.8, "The FedShield-ID Approach", sol_items, title_color=emerald_text)

    # ==================== SLIDE 2: TECHNOLOGY & INNOVATION ====================
    slide2 = prs.slides.add_slide(blank_layout)
    set_background(slide2)
    add_header_footer(slide2, "Slide 2: Technology & Innovation", 2)

    tech_items = [
        ("Backend Framework", "FastAPI microservices coupled with SQLAlchemy ORM running on SQLite/PostgreSQL databases."),
        ("Machine Learning Core", "Scikit-Learn (SGD & Random Forest) for local classification, coupled with linear SHAP explainer code for XAI."),
        ("Interactive SOC Interface", "React 18 frontend built with Tailwind CSS, SVG linkage charts (Knowledge Graph), and live streaming transaction ledgers.")
    ]
    create_panel(slide2, 0.5, 1.7, 5.9, 4.8, "System Architecture & Tech Stack", tech_items, title_color=cyan_text)

    inno_items = [
        ("Post-Quantum KEM Tunnels", "CRYSTALS-Kyber-768 key encapsulation matrix polynomial equations to secure node-to-node communications."),
        ("Differential Privacy (DP)", "Laplacian noise addition to local model parameters controlled by a privacy budget (epsilon) to prevent database leaking."),
        ("Behavioral Biometrics", "Live tracking of keystrokes/minute and mouse jitter standard deviation to stop headless browsers.")
    ]
    create_panel(slide2, 6.9, 1.7, 5.9, 4.8, "Core Cryptography & Telemetry", inno_items, title_color=purple_text)

    # ==================== SLIDE 3: SCALABILITY & DEPLOYMENT ====================
    slide3 = prs.slides.add_slide(blank_layout)
    set_background(slide3)
    add_header_footer(slide3, "Slide 3: Scalability & Deployment", 3)

    feas_items = [
        ("Asynchronous Telemetry Ingestion", "Feeds telemetry data in background threads without interfering with main database queries."),
        ("Sub-Millisecond Overhead", "Crystals-Kyber simulated KEM and local SHAP mathematical computations run in microsecond scales, preserving transaction velocity."),
        ("Containerized Orchestration", "Docker-compose setup makes local deployment reproducible across test nodes and cloud platforms.")
    ]
    create_panel(slide3, 0.5, 1.7, 5.9, 4.8, "Ecosystem Integration & Feasibility", feas_items, title_color=cyan_text)

    imp_items = [
        ("Privacy Compliance", "Meets strict international privacy regulations (GDPR, RBI banking privacy criteria) by removing the centralized data lake."),
        ("Automated Mitigation Playbooks", "Incident response times are reduced using automated GenAI security briefings for SOC operators."),
        ("Money Laundering Prevention", "SVG Knowledge Graph identifies account clusters linked to shared compromised devices or VPN nodes.")
    ]
    create_panel(slide3, 6.9, 1.7, 5.9, 4.8, "Platform Impact & Scale", imp_items, title_color=emerald_text)

    # ==================== SLIDE 4: TEAM & EXECUTION PLAN ====================
    slide4 = prs.slides.add_slide(blank_layout)
    set_background(slide4)
    add_header_footer(slide4, "Slide 4: Team & Execution Plan", 4)

    team_items = [
        ("AI/ML Architect (Mohd. Amaan Hamid)", "Handles federated averages, differential privacy, local ML training models, and SHAP explainability variables."),
        ("Security Engineering", "Configures simulated crystals-kyber KEM, benchmark evaluations, and insider threat audit parameters."),
        ("Web Dashboard Engineer", "Designs React frontend panels, responsive Tailwind styling, and live-streaming ledgers.")
    ]
    create_panel(slide4, 0.5, 1.7, 5.9, 4.8, "Team Roles & Structure", team_items, title_color=purple_text)

    road_items = [
        ("Milestone 1 (Current)", "Completed local prototype running 3 banks with live simulator, SVG knowledge graph, and inline XAI drawer."),
        ("Milestone 2 (Q3 2026)", "Deploy independent containerized bank nodes in clouds and configure multi-channel secure gRPC threads."),
        ("Milestone 3 (Q4 2026)", "Integrate production liboqs library files and package browser-level keyboard/mouse biometrics SDK.")
    ]
    create_panel(slide4, 6.9, 1.7, 5.9, 4.8, "Development Roadmap & Milestones", road_items, title_color=cyan_text)

    output_path = os.path.join(os.getcwd(), "FedShield_ID_Presentation_Slides.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
